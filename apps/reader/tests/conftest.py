import hashlib
import os
import sys
from io import BytesIO
from pathlib import Path

import fitz
import pytest
from docx import Document
from docx.enum.text import WD_BREAK

os.environ.setdefault("READER_SERVICE_TOKEN", "test-service-token")

ROOT = Path(__file__).resolve().parents[1] / "src"
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "HiveForYou Reader PDF fixture page one with enough identity text.")
    second = doc.new_page()
    second.insert_text((72, 72), "Page two continues the synthetic L001-style readable content sample.")
    payload = doc.tobytes()
    doc.close()
    return payload


@pytest.fixture
def sample_pdf_sha256(sample_pdf_bytes: bytes) -> str:
    return hashlib.sha256(sample_pdf_bytes).hexdigest()


@pytest.fixture
def sample_docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("DOCX reproducible page one with sufficient identity text for tests.")
    document.add_paragraph("Second paragraph stays on the same page in reproducible_v1.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def sample_docx_sha256(sample_docx_bytes: bytes) -> str:
    return hashlib.sha256(sample_docx_bytes).hexdigest()


@pytest.fixture
def sample_docx_two_pages_bytes() -> bytes:
    document = Document()
    document.add_paragraph(
        "DOCX page one with sufficient identity text for LibreOffice render tests."
    )
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph(
        "DOCX page two continues with distinct identity text for pagination checks."
    )
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def reader_service_auth_headers() -> dict[str, str]:
    return {"Authorization": "Bearer test-service-token"}


@pytest.fixture
def libreoffice_unpinned(monkeypatch: pytest.MonkeyPatch) -> None:
    from hive_reader.config import settings
    from hive_reader.pipeline.libreoffice_render import (
        LibreOfficeRenderError,
        resolve_libreoffice_binary,
    )

    import hive_reader.pipeline.libreoffice_render as lo_mod

    lo_mod._cached_version = None
    monkeypatch.setattr(settings, "libreoffice_version_pin", "")
    try:
        resolve_libreoffice_binary()
    except LibreOfficeRenderError as exc:
        pytest.skip(str(exc))
